"""Geração de arquivos OFX e Excel."""
from __future__ import annotations

import hashlib
from decimal import Decimal
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .parser import Transaction

BRL = '"R$" #,##0.00;[Red]-"R$" #,##0.00'


def _memo(text: str) -> str:
    """Texto seguro para OFX SGML: sem quebras/controle, escapado e limitado a 255."""
    text = "".join(ch if ch.isprintable() else " " for ch in text)
    return escape(" ".join(text.split())[:255])


def to_ofx(txs: list[Transaction], path: str | Path, account: str = "0000", bank: str = "000") -> None:
    txs = sorted(txs, key=lambda t: t.date)
    bank = (bank or "000").zfill(4)[-4:]
    account = "".join(c for c in (account or "") if c.isalnum() or c == "-")[:22] or "0000"
    seen: dict[str, int] = {}
    now = datetime.now().strftime("%Y%m%d%H%M%S")
    start = txs[0].date.strftime("%Y%m%d") if txs else now[:8]
    end = txs[-1].date.strftime("%Y%m%d") if txs else now[:8]
    total = sum((t.amount for t in txs), Decimal(0))
    rows = []
    for t in txs:
        key = f"{t.date}{t.description}{t.amount}"
        seen[key] = seen.get(key, 0) + 1  # estável entre execuções: reimportar não duplica
        fitid = hashlib.md5(f"{key}{seen[key]}".encode()).hexdigest()[:20]
        rows.append(
            "<STMTTRN>\n"
            f"<TRNTYPE>{t.kind}\n"
            f"<DTPOSTED>{t.date.strftime('%Y%m%d')}\n"
            f"<TRNAMT>{t.amount:.2f}\n"
            f"<FITID>{fitid}\n"
            f"<NAME>{_memo(t.description)[:32]}\n"
            f"<MEMO>{_memo(t.description)}\n"
            "</STMTTRN>"
        )
    body = "\n".join(rows)
    ofx = f"""OFXHEADER:100
DATA:OFXSGML
VERSION:102
SECURITY:NONE
ENCODING:USASCII
CHARSET:1252
COMPRESSION:NONE
OLDFILEUID:NONE
NEWFILEUID:NONE

<OFX>
<SIGNONMSGSRSV1><SONRS>
<STATUS><CODE>0<SEVERITY>INFO</STATUS>
<DTSERVER>{now[:8]}
<LANGUAGE>POR
</SONRS></SIGNONMSGSRSV1>
<BANKMSGSRSV1><STMTTRNRS>
<TRNUID>1
<STATUS><CODE>0<SEVERITY>INFO</STATUS>
<STMTRS>
<CURDEF>BRL
<BANKACCTFROM>
<BANKID>{bank}
<ACCTID>{account}
<ACCTTYPE>CHECKING
</BANKACCTFROM>
<BANKTRANLIST>
<DTSTART>{start}
<DTEND>{end}
{body}
</BANKTRANLIST>
<LEDGERBAL>
<BALAMT>{total:.2f}
<DTASOF>{end}
</LEDGERBAL>
</STMTRS>
</STMTTRNRS></BANKMSGSRSV1>
</OFX>
"""
    Path(path).write_text(ofx, encoding="cp1252", errors="replace")


def to_excel(txs: list[Transaction], path: str | Path) -> None:
    txs = sorted(txs, key=lambda t: t.date)
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"

    navy, light, line = "1F2A44", "EEF2FA", "D5DBE8"
    thin = Side(style="thin", color=line)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:F1")
    ws["A1"] = "Extrato Convertido"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 34

    headers = ["Data", "Descrição", "Tipo", "Crédito", "Débito", "Arquivo"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3B5BDB")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[3].height = 24

    first = 4
    for i, t in enumerate(txs):
        r = first + i
        credit = float(t.amount) if t.amount >= 0 else None
        debit = float(-t.amount) if t.amount < 0 else None
        vals = [t.date, t.description, "Crédito" if credit is not None else "Débito", credit, debit, t.source]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            if i % 2:
                cell.fill = PatternFill("solid", fgColor=light)
        ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=3).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=4).number_format = BRL
        ws.cell(row=r, column=4).font = Font(color="1B7F3B")
        ws.cell(row=r, column=5).number_format = BRL
        ws.cell(row=r, column=5).font = Font(color="C0392B")
    last = first + len(txs) - 1

    tr = last + 2
    ws.cell(row=tr, column=3, value="TOTAIS").font = Font(bold=True)
    for col in (4, 5):
        L = get_column_letter(col)
        cell = ws.cell(row=tr, column=col, value=f"=SUM({L}{first}:{L}{last})")
        cell.number_format = BRL
        cell.font = Font(bold=True)
    ws.cell(row=tr + 1, column=3, value="SALDO").font = Font(bold=True)
    sal = ws.cell(row=tr + 1, column=4, value=f"=D{tr}-E{tr}")
    sal.number_format = BRL
    sal.font = Font(bold=True, color=navy)
    for c in range(3, 6):
        for r in (tr, tr + 1):
            ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="DCE4F7")

    for col, w in zip("ABCDEF", (13, 52, 12, 18, 18, 28)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    if txs:
        ws.auto_filter.ref = f"A3:F{last}"

    # Resumo mensal
    rs = wb.create_sheet("Resumo Mensal")
    for c, h in enumerate(["Mês", "Créditos", "Débitos", "Saldo"], 1):
        cell = rs.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center")
    months: dict[str, list[float]] = {}
    for t in txs:
        m = months.setdefault(t.date.strftime("%Y-%m"), [0.0, 0.0])
        m[0 if t.amount >= 0 else 1] += abs(float(t.amount))
    for i, (k, (cr, db)) in enumerate(sorted(months.items()), 2):
        rs.cell(row=i, column=1, value=k)
        rs.cell(row=i, column=2, value=cr).number_format = BRL
        rs.cell(row=i, column=3, value=db).number_format = BRL
        rs.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = BRL
    for col in "ABCD":
        rs.column_dimensions[col].width = 18
    if months:
        ch = BarChart()
        ch.title = "Créditos x Débitos por mês"
        n = len(months) + 1
        ch.add_data(Reference(rs, min_col=2, max_col=3, min_row=1, max_row=n), titles_from_data=True)
        ch.set_categories(Reference(rs, min_col=1, min_row=2, max_row=n))
        ch.height, ch.width = 8, 16
        rs.add_chart(ch, "F2")

    wb.save(path)
