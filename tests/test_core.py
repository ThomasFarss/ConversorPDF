import sys
from decimal import Decimal as D
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from openpyxl import load_workbook

from core.exporters import to_excel, to_ofx
from core.parser import detect_bank, parse_lines

# layouts simplificados por banco: (linhas, [valores esperados])
CASES = {
    "cresol": (["Cresol", "Extrato 10/2025", "SALDO ANTERIOR 1.000,00 C",
                "02/10/2025 PIX RECEBIDO JOAO 1.500,00 C", "03/10/2025 COMPRA MERCADO 230,45 D",
                "05/10/2025 TARIFA 12,90 D"], ["1500.00", "-230.45", "-12.90"]),
    "sicredi": (["Sicredi 2025", "02/10/2025 PIX RECEBIDO 1.500,00 5.000,00", "03/10/2025 PAGAMENTO BOLETO -230,45 4.769,55"],
                ["1500.00", "-230.45"]),
    "bb": (["Banco do Brasil 10/2025", "02/10/2025 02/10/2025 0000 14397 Pix - Recebido 123 1.500,00 C",
            "03/10/2025 03/10/2025 0000 14397 Pagto fatura 456 230,45 D", "S A L D O 1.000,00 C"],
           ["1500.00", "-230.45"]),
    "santander": (["Santander 2025", "02/10/2025 PIX RECEBIDO 123456 1.500,00 2.500,00",
                   "03/10/2025 COMPRA CARTAO 654321 -230,45 2.269,55"], ["1500.00", "-230.45"]),
    "mercadopago": (["Mercado Pago", "02-10-2025 Pix recebido R$ 1.500,00 R$ 2.500,00",
                     "03-10-2025 Pagamento R$ -230,45 R$ 2.269,55"], ["1500.00", "-230.45"]),
    "nubank": (["Nubank", "Período 01 de OUT de 2025 a 31 de OUT de 2025", "02 OUT 2025 Total de entradas + 1.500,00",
                "Transferência recebida pelo Pix 1.500,00", "JOAO DA SILVA - CPF 123",
                "03 OUT 2025 Total de saídas - 230,45", "Transferência enviada pelo Pix 230,45", "MERCADO LTDA"],
               ["1500.00", "-230.45"]),
    "uniprime": (["Uniprime", "02/10 PIX RECEBIDO 1.500,00C", "03/10 DEB TARIFA 230,45D", "ano 2025"], ["1500.00", "-230.45"]),
    "c6": (["C6 Bank", "02/10/2025 02/10/2025 Entrada Pix recebido 1.500,00", "03/10/2025 03/10/2025 Saída Pix enviado -230,45"],
           ["1500.00", "-230.45"]),
    "sicoob": (["SICOOB 2025", "02/10/2025 PIX RECEBIDO 1.500,00C", "03/10/2025 PAGTO BOLETO 230,45D", "SALDO DO DIA 1.269,55C"],
               ["1500.00", "-230.45"]),
    "btg": (["BTG Pactual", "02/10/2025 Pix recebido R$ 1.500,00", "03/10/2025 Pagamento -R$ 230,45"], ["1500.00", "-230.45"]),
    "itau": (["Itaú 2025", "02/10/2025 PIX TRANSF JOAO 1.500,00", "03/10/2025 SISPAG FORNEC -230,45", "04/10/2025 TARIFA 12,90-"],
             ["1500.00", "-230.45", "-12.90"]),
    "caixa": (["Caixa Econômica Federal", "02/10/2025 123456 CRED PIX 1.500,00 C 2.500,00 C",
               "03/10/2025 654321 DEB COMPRA 230,45 D 2.269,55 C"], ["1500.00", "-230.45"]),
    "inter": (["Banco Inter", "Pix enviado 2025", "2 de outubro de 2025 Saldo do dia: R$ 2.500,00",
               "Pix recebido: JOAO R$ 1.500,00 R$ 2.500,00", "3 de outubro de 2025 Saldo do dia: R$ 2.269,55",
               "Pix enviado: MERCADO -R$ 230,45 R$ 2.269,55"], ["1500.00", "-230.45"]),
}


@pytest.mark.parametrize("name", CASES)
def test_banks(name):
    lines, expected = CASES[name]
    got = parse_lines(lines)
    assert [x.amount for x in got] == [D(v) for v in expected], [(t.date, t.description, t.amount) for t in got]
    assert all(t.date.year == 2025 for t in got)


def test_detect_bank():
    assert detect_bank("Extrato Cresol Cooperativa") == ("Cresol", "133")
    assert detect_bank("Banco do Brasil S.A.")[1] == "001"


def test_export(tmp_path):
    t = parse_lines(CASES["cresol"][0])
    to_ofx(t, tmp_path / "a.ofx", "123-4", "133")
    s = (tmp_path / "a.ofx").read_text(encoding="cp1252")
    assert s.count("<STMTTRN>") == 3 and "<TRNAMT>-230.45" in s and "<BANKID>0133" in s
    to_ofx([], tmp_path / "vazio.ofx")
    to_excel(t, tmp_path / "a.xlsx")
    to_excel([], tmp_path / "vazio.xlsx")
    assert load_workbook(tmp_path / "a.xlsx").sheetnames == ["Lançamentos", "Resumo Mensal"]


def test_descricao_em_outra_linha():
    above = parse_lines(["Cresol 2025", "PIX RECEBIDO JOAO DA SILVA", "02/10/2025 123456 1.500,00 C"])
    assert above[0].description == "PIX RECEBIDO JOAO DA SILVA 123456" or "JOAO" in above[0].description
    below = parse_lines(["Cresol 2025", "02/10/2025 1.500,00 C", "PIX RECEBIDO JOAO DA SILVA"])
    assert below[0].description == "PIX RECEBIDO JOAO DA SILVA"
    s = parse_lines(["02/10/2025 PIX RECEBIDO 123456 1.500,00 C", "JOAO DA SILVA"])
    assert s[0].description == "PIX RECEBIDO 123456 JOAO DA SILVA"
